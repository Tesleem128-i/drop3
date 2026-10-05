"""CSV / Excel export helpers for teacher reports.

A "sheet" is (title, headers, rows). CSV carries one sheet; Excel can carry several.
Cells that start with = + - @ are prefixed with an apostrophe so a student who names themselves
"=HYPERLINK(...)" can't inject a formula into the teacher's spreadsheet.
"""
import csv
import io
import re
from datetime import datetime

from flask import Response

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BAD_SHEET_CHARS = re.compile(r"[\[\]\*\?/\\:]")


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def safe_filename(name):
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._") or "export"
    return name[:80]


def xlsx_available():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _csv_response(filename, headers, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow([_clean(c) for c in r])
    data = "\ufeff" + buf.getvalue()          # BOM so Excel opens UTF-8 names correctly
    return Response(data, mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'})


def _xlsx_response(filename, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for title, headers, rows in sheets:
        base = (_BAD_SHEET_CHARS.sub("", title) or "Sheet")[:28]
        name, n = base, 2
        while name in used:
            name, n = f"{base} {n}", n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append(list(headers))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="6D4AFF")
        for r in rows:
            ws.append([_clean(c) for c in r])
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(_clean(r[i - 1]))) for r in rows[:200] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype=XLSX_MIME,
                    headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'})


def table_response(filename, sheets, fmt="xlsx"):
    """Return a download. Falls back to CSV (first sheet) if Excel was asked for but openpyxl is missing."""
    filename = safe_filename(filename)
    if fmt == "xlsx" and xlsx_available():
        return _xlsx_response(filename, sheets)
    title, headers, rows = sheets[0]
    return _csv_response(filename, headers, rows)